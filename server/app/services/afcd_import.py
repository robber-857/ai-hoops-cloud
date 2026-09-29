"""Strict, offline AFCD Release 3 ingestion. Never infer unknown nutrient values."""

from hashlib import sha256
from io import BytesIO
import json
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re

RELEASE = "AFCD Release 3"
SOURCE_PAGE = "https://www.foodstandards.gov.au/science-data/food-nutrient-databases/afcd/data-files"
LICENCE = "https://www.foodstandards.gov.au/science-data/monitoringnutrients/afcd/datauserlicenceagreement"
PROFILE_SHEET = "All solids & liquids per 100 g"
IDENTITY = ["Public Food Key", "Classification", "Derivation", "Food Name"]
DETAIL_COLUMNS = IDENTITY + [
    "Food Description",
    "Sampling Details",
    "Nitrogen Factor",
    "Fat Factor",
    "Specific Gravity",
    "Analysed Portion",
    "Unanalysed Portion",
]
CORE = {
    "energy_with_fibre_kj": "Energy with dietary fibre, equated (kJ)",
    "protein_g": "Protein (g)",
    "fat_g": "Fat, total (g)",
    "carbohydrate_without_sugar_alcohols_g": "Available carbohydrate, without sugar alcohols (g)",
    "fibre_g": "Total dietary fibre (g)",
    "sugars_g": "Total sugars (g)",
    "sodium_mg": "Sodium (Na) (mg)",
}


class ImportValidationError(ValueError):
    pass


def canonical_bytes(value):
    return json.dumps(
        value,
        sort_keys=True,
        ensure_ascii=False,
        separators=(",", ":"),
        allow_nan=False,
    ).encode("utf-8")


def clean(value):
    return " ".join(str(value).split()) if value is not None else ""


def number(value, location):
    if value is None:
        return None
    if isinstance(value, bool):
        raise ImportValidationError(f"{location}: boolean is not a nutrient")
    try:
        result = Decimal(str(value))
    except InvalidOperation as exc:
        raise ImportValidationError(
            f"{location}: invalid numeric value {value!r}"
        ) from exc
    if not result.is_finite() or result < 0:
        raise ImportValidationError(f"{location}: expected a finite, nonnegative value")
    return format(result, "f")


def table(workbook, sheet, title):
    if sheet not in workbook.sheetnames:
        raise ImportValidationError(f"Missing sheet: {sheet}")
    rows = workbook[sheet].iter_rows(values_only=True)
    if next(rows)[0] != title:
        raise ImportValidationError(f"Unexpected release/title in {sheet}")
    next(rows)
    headers = list(next(rows))
    while headers and headers[-1] is None:
        headers.pop()
    if any(not clean(h) for h in headers):
        raise ImportValidationError(f"Empty header in {sheet}")
    headers = [clean(h) for h in headers]
    if len(headers) != len(set(headers)):
        raise ImportValidationError(f"Duplicate header in {sheet}")
    records = {}
    for line, row in enumerate(rows, 4):
        if all(v is None for v in row):
            continue
        if any(v is not None for v in row[len(headers) :]):
            raise ImportValidationError(f"{sheet}:{line}: value without a header")
        key = row[0]
        if not isinstance(key, str) or not re.fullmatch(r"F\d{6}", key):
            raise ImportValidationError(f"{sheet}:{line}: invalid food key")
        if key in records:
            raise ImportValidationError(f"{sheet}:{line}: duplicate food key {key}")
        records[key] = list(row[: len(headers)])
    return headers, records


def parse_release(source_dir: Path, expected_foods=1588):
    # Read each input once: the recorded hash describes exactly the parsed bytes.
    from openpyxl import load_workbook

    sources, books = {}, {}
    try:
        for stem, official in [
            ("food-details", "Food Details"),
            ("nutrient-profiles", "Nutrient profiles"),
            ("nutrient-details", "Nutrient details"),
        ]:
            raw = (source_dir / f"{stem}.xlsx").read_bytes()
            sources[stem] = {
                "sha256": sha256(raw).hexdigest(),
                "bytes": len(raw),
                "url": "https://www.foodstandards.gov.au/sites/default/files/2025-12/AFCD%20Release%203%20-%20"
                + official.replace(" ", "%20")
                + ".xlsx",
            }
            books[stem] = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        dh, details = table(
            books["food-details"], "Food details", "Release 3 - Food details"
        )
        ph, profiles = table(
            books["nutrient-profiles"],
            PROFILE_SHEET,
            "Release 3 - Nutrient profiles (per 100 g)",
        )
        if dh != DETAIL_COLUMNS or ph[:4] != IDENTITY:
            raise ImportValidationError("Unrecognized food identity/detail columns")
        if set(details) != set(profiles):
            raise ImportValidationError(
                "Food details and nutrient profiles contain different food keys"
            )
        if len(details) != expected_foods or not details:
            raise ImportValidationError(
                f"Expected {expected_foods} foods, found {len(details)}"
            )
        if not set(CORE.values()).issubset(ph[4:]):
            raise ImportValidationError("Required recipe nutrient columns missing")
        nutrients = []
        for header in ph[4:]:
            match = re.fullmatch(r"(.+) \((kJ|g|mg|ug|mg/gN|%T)\)", header)
            if not match:
                raise ImportValidationError(f"Unknown nutrient unit/header: {header}")
            basis = {"%T": "percent_total_fatty_acids", "mg/gN": "per_g_nitrogen"}.get(
                match[2], "per_100g_edible_portion"
            )
            nutrients.append(
                {
                    "key": header,
                    "name": match[1],
                    "unit": match[2],
                    "basis": basis,
                    "official_code": None,
                }
            )
        # Keep definition text and its coordinates, without claiming guessed official codes.
        definitions = {}
        for sheet in [
            "Core nutrients",
            "Proximates",
            "Vitamins",
            "Minerals",
            "Fatty acids",
            "Amino acids",
            "Other",
        ]:
            book = books["nutrient-details"]
            if (
                sheet not in book.sheetnames
                or book[sheet].cell(1, 1).value
                != f"Release 3 - Nutrient details - {sheet}"
            ):
                raise ImportValidationError(
                    f"Unrecognized nutrient definitions: {sheet}"
                )
            definitions[sheet] = {
                cell.coordinate: cell.value
                for row in book[sheet].iter_rows()
                for cell in row
                if cell.value is not None
            }
        foods, warnings = [], []
        for key in sorted(details):
            d, p = details[key], profiles[key]
            if clean(d[3]) != clean(p[3]) or not clean(d[3]):
                raise ImportValidationError(f"{key}: food name mismatch")
            for index in (1, 2):
                if clean(d[index]) != clean(p[index]):
                    warnings.append(
                        {
                            "food_key": key,
                            "field": IDENTITY[index],
                            "food_details": d[index],
                            "nutrient_profiles": p[index],
                        }
                    )
            values = {
                header: number(value, f"{key}/{header}")
                for header, value in zip(ph[4:], p[4:])
            }
            foods.append(
                {
                    "food_key": key,
                    "name": d[3],
                    "details": dict(zip(dh, d)),
                    "profile_identity": dict(zip(IDENTITY, p[:4])),
                    "nutrients": values,
                    # Cooking state is carried by original name/description, never guessed.
                    "cooking_state": None,
                }
            )
        payload = {
            "schema_version": 1,
            "release": RELEASE,
            "basis": "per_100g_edible_portion",
            "source_page": SOURCE_PAGE,
            "sources": sources,
            "attribution": "Food Standards Australia New Zealand (FSANZ), Australian Food Composition Database, Release 3",
            "licence_url": LICENCE,
            "transformation": "Per-100g sheet only; whitespace normalized in nutrient headers; numeric values represented as decimal strings; blanks retained as null.",
            "nutrients": nutrients,
            "core_mapping": CORE,
            "definitions": definitions,
            "foods": foods,
            "warnings": warnings,
        }
        payload["fingerprint"] = sha256(canonical_bytes(payload)).hexdigest()
        return payload
    finally:
        for book in books.values():
            book.close()


def preview(payload):
    values = [v for food in payload["foods"] for v in food["nutrients"].values()]
    return {
        "release": payload["release"],
        "fingerprint": payload["fingerprint"],
        "foods": len(payload["foods"]),
        "nutrients": len(payload["nutrients"]),
        "missing_values": sum(v is None for v in values),
        "zero_values": sum(v is not None and Decimal(v) == 0 for v in values),
        "basis": payload["basis"],
        "sources": payload["sources"],
        "warnings": payload["warnings"],
    }
