import os
from pathlib import Path
from tempfile import TemporaryDirectory
from hashlib import sha256
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
import unittest
from unittest.mock import patch
from contextlib import redirect_stdout, redirect_stderr
from io import StringIO
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, select, func
from sqlalchemy.engine import make_url
from sqlalchemy.orm import Session
from sqlalchemy.exc import IntegrityError
from app.services.afcd_import import (
    CORE,
    DETAIL_COLUMNS,
    IDENTITY,
    PROFILE_SHEET,
    ImportValidationError,
    parse_release,
    preview,
    canonical_bytes,
)
from app.models.food_catalog import FoodRelease, FoodEntry
from app.services.food_catalog_service import import_release


def fixtures(folder):
    book = Workbook()
    sheet = book.active
    sheet.title = "Food details"
    sheet.append(["Release 3 - Food details"])
    sheet.append([])
    sheet.append(DETAIL_COLUMNS)
    sheet.append(
        [
            "F000001",
            "123",
            "Analysed",
            "Synthetic beans, cooked",
            "Cooked edible portion",
            "Synthetic test only",
            6.25,
            1,
            None,
            "100%",
            "0%",
        ]
    )
    book.save(folder / "food-details.xlsx")
    book.close()
    book = Workbook()
    sheet = book.active
    sheet.title = PROFILE_SHEET
    sheet.append(["Release 3 - Nutrient profiles (per 100 g)"])
    sheet.append([])
    sheet.append(
        IDENTITY
        + list(CORE.values())
        + ["Synthetic fatty acid (%T)", "Synthetic amino acid (mg/gN)"]
    )
    sheet.append(
        [
            "F000001",
            123,
            "Analysed",
            "Synthetic beans, cooked",
            418.4,
            10,
            0,
            12.5,
            None,
            1,
            100,
            20,
            50,
        ]
    )
    volume = book.create_sheet("Liquids only per 100 mL")
    volume.append(["Never import me"])
    book.save(folder / "nutrient-profiles.xlsx")
    book.close()
    book = Workbook()
    book.remove(book.active)
    for name in [
        "Core nutrients",
        "Proximates",
        "Vitamins",
        "Minerals",
        "Fatty acids",
        "Amino acids",
        "Other",
    ]:
        sheet = book.create_sheet(name)
        sheet.append([f"Release 3 - Nutrient details - {name}"])
        sheet.append(["Synthetic definition; no official data"])
    book.save(folder / "nutrient-details.xlsx")
    book.close()


def rehash(payload):
    payload["fingerprint"] = sha256(
        canonical_bytes({k: v for k, v in payload.items() if k != "fingerprint"})
    ).hexdigest()
    return payload


class AfcdParserTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.folder = Path(self.temp.name)
        fixtures(self.folder)

    def tearDown(self):
        self.temp.cleanup()

    def edit(self, filename, sheet, cell, value):
        path = self.folder / filename
        book = load_workbook(path)
        book[sheet][cell] = value
        book.save(path)
        book.close()

    def parse(self):
        return parse_release(self.folder, expected_foods=1)

    def test_null_zero_units_traceability_and_deterministic_preview(self):
        payload = self.parse()
        self.assertEqual(payload, self.parse())
        food = payload["foods"][0]
        self.assertIsNone(food["nutrients"][CORE["fibre_g"]])
        self.assertEqual(food["nutrients"][CORE["fat_g"]], "0")
        self.assertEqual(food["nutrients"][CORE["energy_with_fibre_kj"]], "418.4")
        self.assertEqual(payload["nutrients"][-1]["basis"], "per_g_nitrogen")
        self.assertEqual(payload["nutrients"][-2]["basis"], "percent_total_fatty_acids")
        self.assertEqual(preview(payload)["missing_values"], 1)
        self.assertEqual(preview(payload)["zero_values"], 1)
        self.assertIsNone(food["cooking_state"])
        self.assertIn("cooked", food["name"])
        for stem, source in payload["sources"].items():
            self.assertEqual(
                source["sha256"],
                sha256((self.folder / f"{stem}.xlsx").read_bytes()).hexdigest(),
            )

    def test_metadata_disagreement_preserved_as_warning(self):
        self.edit("nutrient-profiles.xlsx", PROFILE_SHEET, "C4", "Recipe")
        payload = self.parse()
        self.assertEqual(payload["warnings"][0]["field"], "Derivation")
        self.assertEqual(
            payload["foods"][0]["profile_identity"]["Derivation"], "Recipe"
        )
        self.assertEqual(payload["foods"][0]["details"]["Derivation"], "Analysed")

    def test_bad_nutrient_values_rejected(self):
        for bad in [-1, "NaN", "Infinity", "trace", True, "=1+2"]:
            with self.subTest(value=bad):
                self.edit("nutrient-profiles.xlsx", PROFILE_SHEET, "E4", bad)
                with self.assertRaises(ImportValidationError):
                    self.parse()

    def test_wrong_name_key_header_release_or_units_rejected(self):
        cases = [
            ("D4", "Wrong food"),
            ("A4", "F000002"),
            ("F3", "Protein (oz)"),
            ("A1", "Release 2"),
            ("F3", CORE["fat_g"]),
            ("A4", "malformed"),
        ]
        for cell, value in cases:
            with self.subTest(cell=cell, value=value):
                fixtures(self.folder)
                self.edit("nutrient-profiles.xlsx", PROFILE_SHEET, cell, value)
                with self.assertRaises(ImportValidationError):
                    self.parse()

    def test_duplicate_food_and_unexpected_count_rejected(self):
        with self.assertRaises(ImportValidationError):
            parse_release(self.folder)
        path = self.folder / "food-details.xlsx"
        book = load_workbook(path)
        sheet = book["Food details"]
        sheet.append([c.value for c in sheet[4]])
        book.save(path)
        book.close()
        with self.assertRaises(ImportValidationError):
            self.parse()

    def test_cli_preview_does_not_connect_and_commit_needs_reviewed_hash(self):
        from app.import_afcd import main

        payload = self.parse()
        with patch("app.import_afcd.parse_release", return_value=payload), patch(
            "sqlalchemy.create_engine"
        ) as connect:
            with patch(
                "sys.argv", ["import_afcd", "--source-dir", str(self.folder)]
            ), redirect_stdout(StringIO()):
                main()
            connect.assert_not_called()
            with patch(
                "sys.argv",
                [
                    "import_afcd",
                    "--source-dir",
                    str(self.folder),
                    "--commit",
                    "--expect-fingerprint",
                    "wrong",
                ],
            ), redirect_stderr(StringIO()):
                with self.assertRaises(SystemExit) as error:
                    main()
                self.assertEqual(error.exception.code, 2)
            connect.assert_not_called()


@unittest.skipUnless(
    os.environ.get("PLAYER_PROFILE_TEST_DATABASE_URL"), "Dedicated PostgreSQL required"
)
class AfcdDatabaseTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        url = make_url(os.environ["PLAYER_PROFILE_TEST_DATABASE_URL"])
        if (
            url.host not in ("localhost", "127.0.0.1")
            or url.database != "ai_hoops_p2_test"
        ):
            raise RuntimeError("Dedicated test database only")
        cls.engine = create_engine(url)

    @classmethod
    def tearDownClass(cls):
        cls.engine.dispose()

    def setUp(self):
        with TemporaryDirectory() as folder:
            fixtures(Path(folder))
            self.payload = parse_release(Path(folder), expected_foods=1)
        self.payload["release"] = f"Synthetic test {uuid4()}"
        rehash(self.payload)

    def save(self):
        with Session(self.engine) as db:
            return import_release(db, self.payload)

    def test_concurrent_import_idempotent_and_content_persisted(self):
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = list(pool.map(lambda _: self.save(), range(2)))
        self.assertEqual(
            sorted(r["status"] for r in results), ["imported", "unchanged"]
        )
        self.assertEqual(results[0]["release_id"], results[1]["release_id"])
        with Session(self.engine) as db:
            food = db.scalar(
                select(FoodEntry).where(
                    FoodEntry.release_id == results[0]["release_id"]
                )
            )
            self.assertEqual(food.content, self.payload["foods"][0])

    def test_changed_release_cannot_overwrite(self):
        saved = self.save()
        self.payload["foods"][0]["name"] = "Modified"
        rehash(self.payload)
        with self.assertRaises(ImportValidationError):
            self.save()
        with Session(self.engine) as db:
            self.assertEqual(
                db.scalar(
                    select(FoodEntry.name).where(
                        FoodEntry.release_id == saved["release_id"]
                    )
                ),
                "Synthetic beans, cooked",
            )

    def test_failed_import_is_atomic_and_tampering_rejected(self):
        self.payload["foods"].append(self.payload["foods"][0])
        with self.assertRaises(ImportValidationError):
            self.save()
        rehash(self.payload)
        with self.assertRaises(IntegrityError):
            self.save()
        with Session(self.engine) as db:
            self.assertEqual(
                db.scalar(
                    select(func.count())
                    .select_from(FoodRelease)
                    .where(FoodRelease.name == self.payload["release"])
                ),
                0,
            )

    def test_existing_catalog_corruption_detected_on_retry(self):
        saved = self.save()
        with Session(self.engine) as db:
            food = db.scalar(
                select(FoodEntry).where(FoodEntry.release_id == saved["release_id"])
            )
            food.content = {**food.content, "name": "Unexpected manual change"}
            db.commit()
        with self.assertRaises(ImportValidationError):
            self.save()
